"""Build native DrawingML charts with embedded XLSX data, never chart screenshots."""
from __future__ import annotations

from io import BytesIO
import math
import os
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
import xlsxwriter

C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART_TYPE = "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PALETTE = ["#2563EB", "#059669", "#D97706", "#7C3AED", "#DC2626", "#0891B2", "#BE185D", "#64748B"]


def chart_workbook(spec: dict) -> tuple[bytes, bytes]:
    """XlsxWriter supplies chart formulas AND matching cached data."""
    stream = BytesIO()
    workbook = xlsxwriter.Workbook(stream, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    sheet = workbook.add_worksheet("Data")
    kind = spec["type"]
    options = {"type": "column" if kind in {"bar", "stacked_bar"} else kind}
    if kind == "stacked_bar":
        options["subtype"] = "stacked"
    chart = workbook.add_chart(options)
    font_name = os.environ.get("DOCUMENT_DOCX_CJK_FONT", "Arial Unicode MS" if os.sys.platform == "darwin" else "Noto Sans CJK SC")
    font = {"name": font_name, "size": 10, "color": "#374151"}
    if kind != "scatter":
        sheet.write_string(0, 0, spec.get("xTitle") or "分类")
        for row, label in enumerate(spec["labels"], 1):
            sheet.write_string(row, 0, label)
    for index, series in enumerate(spec["series"]):
        color = PALETTE[index % len(PALETTE)]
        if kind == "scatter":
            col = index * 2
            sheet.write_string(0, col, series["name"] + " X")
            sheet.write_string(0, col + 1, series["name"])
            for row, point in enumerate(series["points"], 1):
                sheet.write_number(row, col, point["x"])
                sheet.write_number(row, col + 1, point["y"])
            count = len(series["points"])
        else:
            col, count = 0, len(spec["labels"])
            sheet.write_string(0, index + 1, series["name"])
            for row, value in enumerate(series["values"], 1):
                sheet.write_number(row, index + 1, value)
        value_col = col + 1 if kind == "scatter" else index + 1
        settings = {
            "name": ["Data", 0, value_col],
            "categories": ["Data", 1, col, count, col],
            "values": ["Data", 1, value_col, count, value_col],
        }
        if kind in {"pie", "doughnut"}:
            settings["points"] = [{"fill": {"color": PALETTE[i % len(PALETTE)]}, "border": {"color": "white"}} for i in range(count)]
            # Large categorical sets remain available in Edit Data without a wall of labels.
            if count <= 8:
                settings["data_labels"] = {"percentage": True, "position": "best_fit", "font": font}
        elif kind in {"line", "scatter"}:
            settings["line"] = {"color": color, "width": 2} if kind == "line" else {"none": True}
            settings["marker"] = {"type": "circle", "size": 4, "border": {"color": color}, "fill": {"color": color}}
        else:
            settings["fill"] = {"color": color}
            settings["border"] = {"none": True}
        chart.add_series(settings)
    chart.set_title({"none": True})  # Report captions stay searchable, outside the drawing.
    chart.set_chartarea({"border": {"none": True}, "fill": {"color": "white"}})
    chart.set_plotarea({"border": {"none": True}})
    chart.set_legend({"position": "bottom", "font": font})
    if kind not in {"pie", "doughnut"}:
        def axis_title(axis):
            title, unit = spec.get(axis + "Title", ""), spec.get(axis + "Unit", "")
            return f"{title}（{unit}）" if unit and title else title or unit
        x_axis = {"name": axis_title("x"), "name_font": font, "num_font": font, "label_position": "low"}
        if kind != "scatter":
            labels = spec["labels"]
            # Select ticks only; every data point still participates in the plot.
            x_axis["interval_unit"] = max(1, math.ceil(len(labels) / 12))
            if max(map(len, labels)) > 8 or len(labels) > 8:
                x_axis["num_font"] = {**font, "size": 9, "rotation": -45}
        chart.set_x_axis(x_axis)
        y_axis = {"name": axis_title("y"), "name_font": font, "num_font": font, "num_format": "#,##0.##", "major_gridlines": {"visible": True, "line": {"color": "#E5E7EB"}}}
        if kind in {"bar", "stacked_bar"}:
            flat = [v for s in spec["series"] for v in s["values"]]
            if min(flat) >= 0:
                y_axis["min"] = 0
            elif max(flat) <= 0:
                y_axis["max"] = 0
        chart.set_y_axis(y_axis)
    chart.show_blanks_as("gap")
    chart.set_size({"width": 640, "height": 380})
    sheet.insert_chart("T2", chart)
    workbook.close()
    data = stream.getvalue()
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read("xl/charts/chart1.xml")
    return data, xml


def _element(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(key, str(value))
    return node


def embed_native_charts(document_path, charts, markers):
    if not charts:
        return
    document = Document(document_path)
    package = document.part.package
    for spec in charts:
        marker = markers[spec["id"]]
        matches = [p for p in document.paragraphs if p.text == marker]
        if len(matches) != 1:
            raise ValueError(f"图表 {spec['id']} 占位符必须位于正文独立段落，不能放在代码块或表格中。")
        paragraph = matches[0]
        if paragraph.style.name.lower() in {"source code", "verbatim"}:
            raise ValueError("图表占位符不能放在代码块中。")
        workbook, chart_xml = chart_workbook(spec)
        # Allocate package names, including when a reference document contains parts.
        chart_name = package.next_partname("/word/charts/nativeChart%d.xml")
        chart_part = Part(chart_name, CHART_TYPE, b"", package)
        chart_rid = document.part.relate_to(chart_part, RT.CHART)
        workbook_name = package.next_partname("/word/embeddings/chartData%d.xlsx")
        workbook_part = Part(PackURI(workbook_name), XLSX_TYPE, workbook, package)
        data_rid = chart_part.relate_to(workbook_part, RT.PACKAGE)
        tree = etree.fromstring(chart_xml)
        external = etree.Element(f"{{{C}}}externalData", {f"{{{R}}}id": data_rid})
        etree.SubElement(external, f"{{{C}}}autoUpdate", val="0")
        printing = tree.find(f"{{{C}}}printSettings")
        if printing is not None:
            printing.addprevious(external)
        else:
            tree.append(external)
        chart_part._blob = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        # The next section properties describe the section containing this paragraph.
        sections = paragraph._p.xpath("following::w:sectPr")
        section = sections[0] if sections else document.sections[-1]._sectPr
        width = int(section.page_width or Inches(8.27)) - int(section.left_margin or Inches(1)) - int(section.right_margin or Inches(1))
        columns = section.find(qn("w:cols"))
        if columns is not None and int(columns.get(qn("w:num"), "1")) > 1:
            raise ValueError("含图报告暂不支持多栏模板，请使用单栏正文模板。")
        height = int(width * 0.60)
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = bool(spec.get("source") or spec.get("dataDate") or spec.get("caption"))
        title = paragraph.insert_paragraph_before(spec.get("title") or spec["id"], style="Caption")
        title.paragraph_format.keep_with_next = True
        title.paragraph_format.space_before = Pt(8)
        title.paragraph_format.space_after = Pt(4)
        drawing = _element("w:drawing")
        inline = _element("wp:inline", distT=0, distB=0, distL=0, distR=0)
        inline.append(_element("wp:extent", cx=width, cy=height))
        used = [int(v) for v in document._element.xpath("//wp:docPr/@id")]
        inline.append(_element("wp:docPr", id=max(used, default=0) + 1, name=spec["id"], descr=spec.get("title") or spec["id"]))
        graphic = _element("a:graphic")
        graphic_data = _element("a:graphicData", uri=C)
        etree.SubElement(graphic_data, f"{{{C}}}chart", {f"{{{R}}}id": chart_rid})
        graphic.append(graphic_data)
        inline.append(graphic)
        drawing.append(inline)
        paragraph.add_run()._r.append(drawing)
        caption = "\n".join(v for v in [spec.get("caption"), f"来源：{spec['source']}" if spec.get("source") else None, f"数据日期：{spec['dataDate']}" if spec.get("dataDate") else None] if v)
        if caption:
            # Move a new paragraph immediately after the drawing.
            note = document.add_paragraph(caption, style="Caption")
            paragraph._p.addnext(note._p)
            note.paragraph_format.keep_together = True
            note.paragraph_format.keep_with_next = False
            for run in note.runs:
                run.font.size = Pt(9)
    document.save(document_path)
