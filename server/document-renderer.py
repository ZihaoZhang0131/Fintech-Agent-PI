#!/usr/bin/env python3
"""Render restricted Markdown plus generated charts into a DOCX via Pandoc."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from reportlab.graphics import renderPM
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.linecharts import HorizontalLineChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from PIL import Image, ImageDraw, ImageFont
import pypdfium2 as pdfium


def fail(message: str) -> None:
    raise RuntimeError(message)


def configure_cjk_font() -> str:
    regular_candidates = [
        os.environ.get("DOCUMENT_CJK_FONT_PATH"),
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    ]
    bold_candidates = [
        os.environ.get("DOCUMENT_CJK_BOLD_FONT_PATH"),
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc",
    ]
    regular_path = next((Path(value) for value in regular_candidates if value and Path(value).is_file()), None)
    bold_path = next((Path(value) for value in bold_candidates if value and Path(value).is_file()), regular_path)
    if not regular_path or not bold_path:
        fail("未找到可嵌入 PDF 的中文字体。请通过 DOCUMENT_CJK_FONT_PATH 配置 TTF/TTC 字体。")
    regular_name, bold_name = "EmbeddedCjk", "EmbeddedCjkBold"
    if regular_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(regular_name, str(regular_path)))
    if bold_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(bold_name, str(bold_path)))
    pdfmetrics.registerFontFamily(
        regular_name,
        normal=regular_name,
        bold=bold_name,
        italic=regular_name,
        boldItalic=bold_name,
    )
    return regular_name


def render_chart(chart: dict, directory: Path) -> tuple[str, Path]:
    chart_id = chart["id"]
    labels = chart["labels"]
    series = chart["series"]
    if any(len(item["values"]) != len(labels) for item in series):
        fail(f"图表 {chart_id} 的每个序列必须与 labels 等长。")

    font = configure_cjk_font()
    drawing = Drawing(760, 420)
    title = chart.get("title") or chart_id
    drawing.add(String(24, 392, title, fontName=font, fontSize=16, fillColor=colors.HexColor("#1f2937")))
    values = [item["values"] for item in series]
    if chart["type"] == "bar":
        graph = VerticalBarChart()
        graph.x, graph.y, graph.width, graph.height = 60, 70, 650, 270
        graph.data = values
        graph.categoryAxis.categoryNames = labels
        graph.categoryAxis.labels.fontName = font
        graph.categoryAxis.labels.fontSize = 8
        graph.valueAxis.labels.fontName = font
        graph.valueAxis.labels.fontSize = 8
        graph.barSpacing = 5
        graph.groupSpacing = 12
        for index in range(len(values)):
            graph.bars[index].fillColor = [colors.HexColor("#2563eb"), colors.HexColor("#059669"), colors.HexColor("#d97706"), colors.HexColor("#7c3aed")][index % 4]
    else:
        graph = HorizontalLineChart()
        graph.x, graph.y, graph.width, graph.height = 60, 70, 650, 270
        graph.data = values
        graph.categoryAxis.categoryNames = labels
        graph.categoryAxis.labels.fontName = font
        graph.categoryAxis.labels.fontSize = 8
        graph.valueAxis.labels.fontName = font
        graph.valueAxis.labels.fontSize = 8
        graph.joinedLines = 1
        for index in range(len(values)):
            graph.lines[index].strokeColor = [colors.HexColor("#2563eb"), colors.HexColor("#059669"), colors.HexColor("#d97706"), colors.HexColor("#7c3aed")][index % 4]
            graph.lines[index].strokeWidth = 2
    drawing.add(graph)
    for index, item in enumerate(series):
        x = 65 + index * 160
        drawing.add(String(x, 36, f"{item['name']}", fontName=font, fontSize=9, fillColor=colors.HexColor("#374151")))
    target = directory / f"chart-{chart_id}.png"
    try:
        renderPM.drawToFile(drawing, str(target), "PNG", dpi=144)
    except Exception:
        render_chart_with_pillow(chart, target)
    return chart_id, target


def pillow_font(size: int):
    for candidate in ["/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Light.ttc"]:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default()


def render_chart_with_pillow(chart: dict, target: Path) -> None:
    width, height = 1520, 840
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font, label_font = pillow_font(32), pillow_font(18)
    colors_list = ["#2563eb", "#059669", "#d97706", "#7c3aed"]
    left, top, right, bottom = 130, 120, 1420, 650
    values = [value for item in chart["series"] for value in item["values"]]
    minimum, maximum = min(0, min(values)), max(0, max(values))
    if minimum == maximum:
        maximum = minimum + 1
    span = maximum - minimum
    draw.text((48, 40), chart.get("title") or chart["id"], font=title_font, fill="#1f2937")
    draw.line((left, top, left, bottom, right, bottom), fill="#6b7280", width=2)
    for step in range(5):
        y = top + (bottom - top) * step / 4
        value = maximum - span * step / 4
        draw.line((left, y, right, y), fill="#e5e7eb", width=1)
        draw.text((18, y - 10), f"{value:.2g}", font=label_font, fill="#4b5563")
    labels = chart["labels"]
    group_width = (right - left) / len(labels)
    for index, label in enumerate(labels):
        x = left + group_width * (index + 0.5)
        draw.text((x - 20, bottom + 18), label, font=label_font, fill="#4b5563")
    if chart["type"] == "bar":
        bar_width = group_width / (len(chart["series"]) + 2)
        for series_index, series in enumerate(chart["series"]):
            for index, value in enumerate(series["values"]):
                x1 = left + group_width * index + bar_width * (series_index + 1)
                x2 = x1 + bar_width * 0.8
                y = bottom - (value - minimum) / span * (bottom - top)
                draw.rectangle((x1, y, x2, bottom - (-minimum / span * (bottom - top))), fill=colors_list[series_index % len(colors_list)])
    else:
        for series_index, series in enumerate(chart["series"]):
            points = []
            for index, value in enumerate(series["values"]):
                x = left + group_width * (index + 0.5)
                y = bottom - (value - minimum) / span * (bottom - top)
                points.append((x, y))
            draw.line(points, fill=colors_list[series_index % len(colors_list)], width=5, joint="curve")
            for point in points:
                draw.ellipse((point[0] - 5, point[1] - 5, point[0] + 5, point[1] + 5), fill=colors_list[series_index % len(colors_list)])
    for index, series in enumerate(chart["series"]):
        x = left + index * 200
        draw.rectangle((x, 730, x + 22, 752), fill=colors_list[index % len(colors_list)])
        draw.text((x + 32, 730), series["name"], font=label_font, fill="#374151")
    image.save(target, "PNG")


def replace_chart_markers(markdown: str, charts: list[dict], directory: Path) -> str:
    generated: dict[str, Path] = {}
    for chart in charts:
        chart_id, target = render_chart(chart, directory)
        if chart_id in generated:
            fail(f"图表 ID 重复：{chart_id}。")
        generated[chart_id] = target
    for chart_id, target in generated.items():
        marker = "{{chart:" + chart_id + "}}"
        if markdown.count(marker) != 1:
            fail(f"图表 {chart_id} 必须在 Markdown 中恰好使用一次 {marker} 占位符。")
        markdown = markdown.replace(marker, f"![{chart_id}]({target.name})")
    if "{{chart:" in markdown:
        fail("Markdown 包含未定义的图表占位符。")
    allowed_assets = {target.name for target in generated.values()}
    for match in re.finditer(r"!\[[^\]]*\]\(([^)\s]+)(?:\s+[^)]*)?\)", markdown):
        asset = match.group(1).strip("<>")
        if asset not in allowed_assets:
            fail("Markdown 图片只能引用本次 charts 参数生成的图表，不能引用远程或本机文件。")
    return markdown


def apply_default_cjk_styles(document_path: Path) -> None:
    document = Document(str(document_path))
    font_name = os.environ.get("DOCUMENT_DOCX_CJK_FONT") or ("Arial Unicode MS" if sys.platform == "darwin" else "Noto Sans CJK SC")
    for style_name in ["Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    for paragraph in list(document.paragraphs) + [paragraph for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]:
        for run in paragraph.runs:
            run.font.name = font_name
            run._element.rPr.rFonts.set(qn("w:ascii"), font_name)
            run._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    document.save(str(document_path))


def inline_pdf_text(value: str) -> str:
    escaped = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    escaped = re.sub(r"`([^`]+)`", r"\1", escaped)
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escaped)


def render_pdf_from_markdown(markdown: str, directory: Path, output: Path, render_directory: Path) -> tuple[int, int]:
    """Create a CJK-safe PDF and render every page for deterministic validation."""
    font = configure_cjk_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("CjkBody", parent=styles["BodyText"], fontName=font, fontSize=10.5, leading=18, spaceAfter=7)
    headings = {level: ParagraphStyle(f"CjkHeading{level}", parent=styles[f"Heading{level}"], fontName=font, fontSize={1: 22, 2: 16, 3: 13}[level], leading={1: 30, 2: 24, 3: 20}[level], textColor=colors.HexColor("#1f2937"), spaceBefore=10 if level > 1 else 0, spaceAfter=8, keepWithNext=True) for level in (1, 2, 3)}
    bullet = ParagraphStyle("CjkBullet", parent=body, leftIndent=16, firstLineIndent=-10, bulletIndent=2)
    numbered = ParagraphStyle("CjkNumbered", parent=body, leftIndent=20, firstLineIndent=-16, bulletIndent=2)
    quote = ParagraphStyle(
        "CjkQuote",
        parent=body,
        leftIndent=10,
        rightIndent=10,
        borderWidth=0.5,
        borderColor=colors.HexColor("#d1d5db"),
        borderPadding=7,
        backColor=colors.HexColor("#f7f8fa"),
        textColor=colors.HexColor("#4b5563"),
        spaceAfter=5,
    )
    story, lines, index = [], markdown.splitlines(), 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if re.fullmatch(r"-{3,}", line):
            story.append(Spacer(1, 5))
            index += 1
            continue
        heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading_match:
            story.append(Paragraph(inline_pdf_text(heading_match.group(2)), headings[len(heading_match.group(1))]))
            index += 1
            continue
        image_match = re.fullmatch(r"!\[[^\]]*\]\(([^)]+)\)", line)
        if image_match:
            image = PdfImage(str(directory / image_match.group(1)))
            image._restrictSize(170 * mm, 105 * mm)
            story.extend([Spacer(1, 4), image, Spacer(1, 10)])
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    rows.append([Paragraph(inline_pdf_text(cell), body) for cell in cells])
                index += 1
            if rows:
                column_count = max(len(row) for row in rows)
                width_patterns = {
                    2: [0.30, 0.70],
                    3: [0.24, 0.46, 0.30],
                    4: [0.16, 0.28, 0.28, 0.28],
                }
                ratios = width_patterns.get(column_count, [1 / column_count] * column_count)
                table = Table(rows, repeatRows=1, hAlign="LEFT", colWidths=[166 * mm * ratio for ratio in ratios])
                table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), font), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5e7eb")), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d1d5db")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
                story.extend([table, Spacer(1, 10)])
            continue
        if line.startswith(("- ", "* ")):
            story.append(Paragraph(inline_pdf_text(line[2:]), bullet, bulletText="•"))
            index += 1
            continue
        numbered_match = re.match(r"^(\d+)\.\s+(.+)$", line)
        if numbered_match:
            story.append(Paragraph(inline_pdf_text(numbered_match.group(2)), numbered, bulletText=f"{numbered_match.group(1)}."))
            index += 1
            continue
        if line.startswith(">"):
            quote_lines = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(lines[index].strip()[1:].strip())
                index += 1
            story.append(Paragraph("<br/>".join(inline_pdf_text(item) for item in quote_lines), quote))
            continue
        paragraph = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not re.match(r"^(#{1,3})\s+|^\||^[-*] |^\d+\.\s+|^>|^-{3,}$", lines[index].strip()):
            paragraph.append(lines[index].strip())
            index += 1
        story.append(Paragraph(inline_pdf_text(" ".join(paragraph)), body))
    document = SimpleDocTemplate(str(output), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm)

    def draw_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#9ca3af"))
        canvas.setFont(font, 8)
        canvas.drawCentredString(A4[0] / 2, 10 * mm, str(doc.page))
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)
    render_directory.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(output))
    try:
        page_count = len(pdf)
        if page_count < 1:
            fail("生成的 PDF 没有有效页面。")
        for index in range(page_count):
            page = pdf[index]
            bitmap = page.render(scale=1.6)
            try:
                bitmap.to_pil().save(render_directory / f"page-{index + 1}.png", format="PNG")
            finally:
                bitmap.close()
                page.close()
    finally:
        pdf.close()
    return page_count, page_count


def main() -> None:
    if len(sys.argv) not in {3, 5, 7} or (len(sys.argv) >= 5 and sys.argv[3] != "--pdf") or (len(sys.argv) == 7 and sys.argv[5] != "--pdf-render-dir"):
        fail("用法：document-renderer.py 输入 JSON 输出 DOCX [--pdf 输出 PDF [--pdf-render-dir 逐页预览目录]]")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    pdf_output = Path(sys.argv[4]).resolve() if len(sys.argv) >= 5 else None
    pdf_render_directory = Path(sys.argv[6]).resolve() if len(sys.argv) == 7 else None
    workdir = output.parent
    payload = json.loads(source.read_text(encoding="utf-8"))
    markdown = replace_chart_markers(payload["markdown"], payload.get("charts") or [], workdir)
    markdown_path = workdir / "document.md"
    markdown_path.write_text(markdown, encoding="utf-8")
    pandoc = os.environ.get("PANDOC_PATH")
    if not pandoc or not Path(pandoc).is_file():
        fail("未安装受管 Pandoc。请先运行 npm run documents:setup。")
    args = [
        pandoc,
        str(markdown_path),
        "--from=gfm+pipe_tables+task_lists",
        "--to=docx",
        "--standalone",
        f"--resource-path={workdir}",
        "--output",
        str(output),
    ]
    reference = payload.get("referenceDocx")
    if reference:
        args.extend(["--reference-doc", reference])
    completed = subprocess.run(args, cwd=workdir, text=True, capture_output=True, check=False)
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout or "未知错误").strip()
        fail(f"Pandoc 生成 DOCX 失败：{detail[:2000]}")
    if not output.is_file() or output.stat().st_size < 4:
        fail("Pandoc 未生成有效 DOCX 文件。")
    if not reference:
        apply_default_cjk_styles(output)
    if pdf_output:
        render_directory = pdf_render_directory or pdf_output.parent / f"{pdf_output.stem}-pages"
        page_count, rendered_pages = render_pdf_from_markdown(markdown, workdir, pdf_output, render_directory)
        print("DOCUMENT_RESULT " + json.dumps({"pageCount": page_count, "renderedPages": rendered_pages}))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
